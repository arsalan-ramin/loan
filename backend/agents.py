"""Evidence-based agents with Anthropic conversation and decision explanations."""

import csv
import hashlib
import io
import json
import logging
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader

from .models import (Document, Finding, PolicyEvaluation, PolicyProfileResult,
                     PolicyRuleResult, Report, Session, default_rules)
from .llm import AnthropicLLM, LLMError
from .storage import Store
from .excel import extract_excel
from .underwriting_engine import evaluate_underwriting
from .underwriting_models import SOURCE_DOCUMENT, SOURCE_SHA256, UnderwritingBatch

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xlsx"}
MAX_CONTENT = 2_000_000
LOGGER = logging.getLogger(__name__)


class IntakeAgent:
    def check(self, session: Session) -> str:
        if not session.documents:
            raise ValueError("Upload at least one document before processing.")
        if not session.policy_rules and not any(rule.enabled for rule in session.rules):
            raise ValueError("Enable at least one rule before processing.")
        if session.policy_rules:
            return f"Validated {len(session.documents)} documents. Draft underwriting safeguards apply; optional document checks are separate."
        return f"Validated {len(session.documents)} documents and {sum(rule.enabled for rule in session.rules)} active document checks."

    def greeting(self) -> str:
        return ("Hello! I’m your document review assistant. First, upload the files you want me to review "
            "(XLSX, PDF, DOCX, TXT, Markdown, CSV, or JSON). For MTB reviews, upload Query 1 (Loans, Accounts & Financials), "
            "Query 2 (SBSS & Entity Relationships), and Query 4 (Credit Reports). These exports link through explicit Account ID columns; "
            "the app does not join by row position or borrower name. Then validate the extracted attributes and select Process documents. "
                "I’ll extract content and metadata, evaluate each file, and prepare a downloadable HTML decision report. "
                "Your conversation stays saved in this workspace.")


class ExtractionAgent:
    def extract(self, document: Document, path: Path, *, include_text: bool = True) -> None:
        # Hash XLSX incrementally and pass its path to the read-only parser.
        # Other formats retain their existing byte-based parsers.
        if document.extension == ".xlsx":
            digest = hashlib.sha256()
            with path.open("rb") as source:
                for chunk in iter(lambda: source.read(1024 * 1024), b""):
                    digest.update(chunk)
            raw = None
            size_bytes, sha256 = path.stat().st_size, digest.hexdigest()
        else:
            raw = path.read_bytes()
            size_bytes, sha256 = len(raw), hashlib.sha256(raw).hexdigest()
        metadata: dict = {
            "filename": document.name, "extension": document.extension,
            "size_bytes": size_bytes, "sha256": sha256,
            "extracted_at": datetime.now(timezone.utc).isoformat(),
        }
        if document.extension == ".xlsx":
            content, excel_metadata = extract_excel(path, include_text=include_text)
            metadata.update(excel_metadata)
        elif document.extension == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError(f"{document.name}: encrypted PDFs are not supported.")
            if len(reader.pages) > 500:
                raise ValueError(f"{document.name}: PDFs are limited to 500 pages.")
            parts = []
            total = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_CONTENT:
                    raise ValueError(f"{document.name}: extracted text exceeds the 2 million character limit.")
                parts.append(text)
            content = "\n\n".join(parts)
            metadata["pages"] = len(reader.pages)
            if reader.metadata:
                for key in ("title", "author", "subject", "creator"):
                    value = getattr(reader.metadata, key, None)
                    if value:
                        metadata[key] = str(value)
        elif document.extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError(f"{document.name}: expanded DOCX is too large.")
            word = WordDocument(io.BytesIO(raw))
            paragraphs = [p.text for p in word.paragraphs]
            tables = [" | ".join(cell.text for cell in row.cells) for table in word.tables for row in table.rows]
            content = "\n".join(paragraphs + tables)
            for key in ("title", "author", "subject"):
                value = getattr(word.core_properties, key, None)
                if value:
                    metadata[key] = str(value)
            metadata["paragraphs"] = len(paragraphs)
        else:
            try:
                content = raw.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ValueError(f"{document.name}: please save text files as UTF-8.") from exc
            if "\x00" in content:
                raise ValueError(f"{document.name}: binary content is not supported.")
            if document.extension == ".json":
                value = json.loads(content)
                metadata["json_type"] = type(value).__name__
                if isinstance(value, dict):
                    metadata["json_keys"] = list(value.keys())[:100]
                # Retain the original JSON text for literal, auditable rule matching.
            elif document.extension == ".csv":
                rows = csv.reader(io.StringIO(content))
                metadata["columns"] = next(rows, [])
                metadata["rows"] = sum(1 for _ in rows)
        # XLSX has its own larger bounded limit in extract_excel; keep the
        # existing 2-million-character ceiling for all other formats.
        if document.extension != ".xlsx" and len(content) > MAX_CONTENT:
            raise ValueError(f"{document.name}: extracted text exceeds the 2 million character limit.")
        metadata["word_count"] = len(content.split())
        metadata["character_count"] = len(content)
        if include_text and not content.strip():
            metadata["warning"] = "No readable text found. Scanned documents require OCR before upload."
        document.content = content
        document.metadata = metadata


class RulesAgent:
    def evaluate(self, session: Session) -> list[Finding]:
        findings = []
        is_underwriting = bool(session.policy_rules)
        if is_underwriting:
            findings.append(Finding(rule_id="system-mtb-guideline", rule_name="MTB draft policy requires validation", status="review",
                                    detail="The supplied design guideline is a draft, not approved credit policy. Page-cited policy routes require authorized review; generic checks are not lending decisions.",
                                    document="MTB dataset", evidence=f"{SOURCE_DOCUMENT}: verified v4 draft diagnostics, not operational lending authority."))
        for document in session.documents:
            if document.extension == ".xlsx" and document.metadata.get("data_row_count", 0) == 0:
                findings.append(Finding(rule_id="system-excel-records", rule_name="Excel data records required", status="review",
                                        detail="The workbook contains headers but no data records outside reference sheets. No borrower decision can be made.",
                                        document=document.name, evidence="metadata.data_row_count = 0"))
            # Empty extraction always requires attention, even if all keyword rules happen to pass.
            if not document.content.strip():
                findings.append(Finding(rule_id="system-readable", rule_name="Readable content", status="review",
                                        detail="No readable text; OCR or a text-based file is required.",
                                        document=document.name, evidence="Extracted text is empty."))
            for rule in session.rules:
                if is_underwriting and rule in default_rules():
                    continue
                if not rule.enabled:
                    continue
                content, target = document.content, rule.value
                location = content.casefold().find(target.casefold())
                if rule.kind == "required_text":
                    passed = location >= 0
                    detail = f'Required phrase “{target}” {"found" if passed else "not found"}.'
                    evidence = content[max(0, location - 70):location + len(target) + 130] if passed else "No matching text in the extracted content."
                elif rule.kind == "forbidden_text":
                    passed = location < 0
                    detail = f'Forbidden phrase “{target}” {"not found" if passed else "found"}.'
                    evidence = "No matching text in the extracted content." if passed else content[max(0, location - 70):location + len(target) + 130]
                elif rule.kind == "min_words":
                    count = int(document.metadata.get("word_count", 0))
                    passed = count >= int(target)
                    detail = f"{count} words extracted; at least {target} required."
                    evidence = f"metadata.word_count = {count}"
                else:
                    value = document.metadata.get(target)
                    passed = value is not None and bool(str(value).strip())
                    detail = f'Metadata field “{target}” {"present" if passed else "missing"}.'
                    evidence = f"{target} = {value!s}"[:500]
                findings.append(Finding(rule_id=rule.id, rule_name=rule.name, status="pass" if passed else rule.severity,
                                        detail=detail, document=document.name, evidence=evidence))
        return findings


def underwriting_policy_summary(batch: UnderwritingBatch, session: Session) -> PolicyEvaluation:
    """Read-only UI adapter from v4 results, NEVER a second policy execution.

    The legacy statuses mean attention/not-attention/unknown here, not lending
    outcomes. Profiles represent explicit loan IDs only, never arbitrary rows.
    Non-policy diagnostics have no PDF page; null explicitly denotes this.
    """
    missing = list(batch.data_quality_issues)
    missing.extend(f"{doc.name}: no data records outside reference sheets."
                   for doc in session.documents
                   if doc.extension == ".xlsx" and not doc.metadata.get("records"))
    for loan in batch.loan_results:
        missing.extend(f"{loan.loan_id}: {item}" for item in loan.missing_inputs + loan.conflicts)
    return PolicyEvaluation(
        policy_version=batch.policy_version, source_document=batch.source_document,
        source_sha256=SOURCE_SHA256 if batch.source_available else None,
        recommended_route="MANUAL_REVIEW", missing_information=list(dict.fromkeys(missing)),
        profiles=[PolicyProfileResult(profile_id=loan.loan_id, route=loan.recommendation_path,
                                      missing_information=loan.missing_inputs + loan.conflicts,
                                      evidence=[item.id for item in loan.source_evidence])
                  for loan in batch.loan_results],
        rule_results=[PolicyRuleResult(
            rule_id=rule.rule_id,
                status=("triggered" if rule.status == "FAIL" else
                    "unknown" if rule.status == "UNKNOWN" else "not_triggered"),
            route=loan.recommendation_path, page=rule.page,
            detail=f"Loan {loan.loan_id} · {rule.status}: {rule.detail}", evidence=rule.evidence_refs,
        ) for loan in batch.loan_results for rule in loan.rule_results],
        limitations=[
            "Compatibility summary of v4 only; legacy draft rules were not executed. "
            "A null page means a data diagnostic without a PDF citation; NOT_APPLICABLE is not a passed policy gate.",
            "Verified source integrity is not approved lending authority. No lending approval or denial.",
            "Per-loan explanations are deterministic engine text; Claude provides bounded aggregate commentary only.",
        ],
    )


class DecisionAgent:
    def decide(self, session: Session, findings: list[Finding]) -> Report:
        failed = sum(f.status == "fail" for f in findings)
        reviews = sum(f.status == "review" for f in findings)
        decision = "rejected" if failed else "needs_review" if reviews else "approved"
        report = Report(decision=decision, findings=findings, document_count=len(session.documents),
                      summary=f"Reviewed {len(session.documents)} document(s) against {sum(rule.enabled for rule in session.rules)} active rules. "
                              f"{sum(f.status == 'pass' for f in findings)} checks passed, {reviews} need review, and {failed} failed. "
                              "This is a rule-based assessment, not a legal or professional determination.")
        if session.policy_rules:
            report.underwriting = evaluate_underwriting(
                session.documents, session.mode,
                selected_account_id=session.selected_account_id,
            )
            report.policy_evaluation = underwriting_policy_summary(report.underwriting, session)
            report.decision = "needs_review"
            report.summary = (f"Document checks completed for {len(session.documents)} file(s). "
                              f"V4 {session.mode} underwriting: {len(report.underwriting.loan_results)} explicit loan(s); "
                              "aggregate route MANUAL_REVIEW. "
                              "No lending approval or denial. Policy-owner validation and authorized human review are required.")
            LOGGER.info("underwriting_evaluation mode=%s version=%s loans=%d rules=%d issues=%d",
                        session.mode, report.underwriting.policy_version, len(report.underwriting.loan_results),
                        sum(len(loan.rule_results) for loan in report.underwriting.loan_results),
                        len(report.underwriting.data_quality_issues))
        llm = AnthropicLLM()
        try:
            report.llm_analysis = llm.explain(report)
            report.llm_model = llm.model
        except LLMError as exc:
            # Do not discard valid deterministic findings when the provider is unavailable.
            report.llm_error = str(exc)
        return report


class MemoryAgent:
    def __init__(self, store: Store):
        self.store = store

    def index(self, session: Session) -> str:
        self.store.forget_documents(session.id)
        for doc in session.documents:
            for index, start in enumerate(range(0, len(doc.content), 1200)):
                self.store.remember(session.id, f"doc:{doc.id}:{index}", doc.content[start:start + 1400], "document", doc.name)
        for message in session.messages:
            self.store.remember(session.id, message.id, message.content, "conversation", message.role)
        return "Document evidence and conversation saved to local ChromaDB."

    def respond(self, session: Session, query: str) -> str:
        return AnthropicLLM().chat(self.store, session, query)