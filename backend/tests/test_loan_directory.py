"""Synthetic-only explicit loan directory safety; providers never make live calls."""

import io
import sqlite3
from collections import defaultdict
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import backend.underwriting_engine as engine
from backend.assistant_workflow import ACCOUNT_INDEX_VERSION, _discover_accounts
from backend.entity_graph import build_entity_graph
from backend.llm import AnthropicLLM, LLMError
from backend.main import create_app
from backend.models import AccountCandidate, Document, LoanCandidate


JUNCTION = "SCOC__LCRDataSourceAnalysis__r.SCOC__LCREntity__r.SCOC__Entity_Involvement__r."


def workbook(records):
    groups = defaultdict(list)
    for sheet, values in records:
        groups[sheet].append(values)
    book = Workbook()
    book.remove(book.active)
    for name, rows in groups.items():
        sheet = book.create_sheet(name)
        headers = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key) for key in headers])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    provider = Mock(side_effect=LLMError("Synthetic offline provider stub"))
    monkeypatch.setattr(AnthropicLLM, "complete", provider)
    with TestClient(create_app(tmp_path / "data")) as active:
        active.app.state.provider_stub = provider
        yield active


def upload_guarantors(client):
    sid = client.post("/api/sessions", json={}).json()["id"]
    q1 = [("Sheet1", {
        "Id": f"ROW-{loan}-{account}", JUNCTION + "LLC_BI__Loan__c": loan,
        JUNCTION + "LLC_BI__Account__c": account,
        JUNCTION + "LLC_BI__Borrower_Type__c": "Guarantor",
        JUNCTION + "LLC_BI__Account__r.Name": "Synthetic guarantor " + account,
        "as_of_date": "2026-01-01",
        "SCOC__CharacteristicKey__c": "numCollection", "SCOC__CharacteristicValue__c": "0",
    }) for loan, account in (("L1", "G1"), ("L1", "G2"), ("L2", "G3"), ("L3", "G1"))]
    q2 = [(sheet, {**{key: value for key, value in values.items()
                            if not key.startswith("SCOC__Characteristic")},
                   "Id": "IND-" + values["Id"],
                   "SCOC__IndicatorKey__c": "thinFileIndicator", "SCOC__IndicatorValue__c": "N"})
          for sheet, values in q1]
    q4 = [("CreditReport", {"Id": "R-" + account, "account_id": account,
                            "report_date": "2025-12-01", "status": "Completed",
                            "personal_fico": "740", "report_category": "personal"})
          for account in ("G1", "G2", "G3", "UNRELATED")]
    response = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", (f"MTB_Query_{query}.xlsx", workbook(records)))
        for query, records in ((1, q1), (2, q2), (4, q4))
    ])
    assert response.status_code == 200
    return sid


def test_guarantor_only_directory_selection_revalidation_and_confirmed_process(client, monkeypatch):
    sid = upload_guarantors(client)
    base = f"/api/sessions/{sid}"
    validated = client.post(base + "/preflight")
    assert validated.status_code == 200
    result = validated.json()["preflight"]
    assert result["requires_loan_number"] is True
    assert result["loan_count"] == 0 and result["account_count"] == 0
    assert "loan_candidates" not in result and "account_candidates" not in result
    assert "What loan number" in result["suggested_next_step"]
    assert client.get(base + "/accounts").json()["results"] == []
    first = client.get(base + "/loans", params={"limit": 2}).json()
    assert first["total"] == 0
    assert first["results"] == []
    assert client.get(base + "/loans?q=G1").json()["total"] == 0
    assert client.put(base + "/loan", json={"loan_id": "ROW-L1-G1"}).status_code == 404
    assert client.put(base + "/account", json={"account_id": "G1"}).status_code == 404
    assert client.post(base + "/process").status_code == 409
    selected = client.put(base + "/loan", json={"loan_id": "L1"})
    assert selected.status_code == 200
    assert selected.json()["selected_account_id"] is None
    assert "No unique verified borrower" in selected.json()["messages"][-1]["content"]
    again = client.post(base + "/preflight").json()
    assert again["selected_loan_id"] == "L1" and again["selected_account_id"] is None
    client.app.state.provider_stub.assert_not_called()

    report_selector = Mock(wraps=engine._select_report)
    graph_builder = Mock(wraps=engine.build_entity_graph)
    monkeypatch.setattr(engine, "_select_report", report_selector)
    monkeypatch.setattr(engine, "build_entity_graph", graph_builder)
    response = client.post(base + "/assistant/process", json={
        "confirmed": True, "expected_updated_at": again["updated_at"],
    })
    assert response.status_code == 202
    client.app.state.executor.shutdown(wait=True)
    completed = client.get(base).json()
    assert completed["status"] == "completed", completed["error"]
    loans = completed["report"]["underwriting"]["loan_results"]
    assert [loan["loan_id"] for loan in loans] == ["L1"]
    loan = loans[0]
    assert loan["borrower_account_id"] is None
    assert "explicit_unique_borrower_account_relationship" in loan["missing_inputs"]
    assert "borrower_personal_fico_for_screenshot_average" in loan["missing_inputs"]
    assert loan["final_action"] == "NO_OPERATIONAL_ACTION"
    assert loan["authorization"] == "NOT_ISSUED"
    # No account filter before full graph validation, then selected-loan scoping.
    assert graph_builder.call_args.kwargs["selected_account_id"] is None
    assert {call.args[1] for call in report_selector.call_args_list} == {"G1", "G2"}
    assert all(call.args[0].table("CreditReport")[report_id].fields["account_id"] == account
               for call in report_selector.call_args_list
               for account in [call.args[1]] for report_id in ["R-" + account])
    assert client.get(base + "/underwriting").status_code == 200
    assert client.post(base + "/assistant/inspect", json={"loan_id": "L1", "topic": "blockers"}).status_code == 200


@pytest.mark.parametrize("endpoint", ["accounts", "loans", "loan"])
def test_legacy_refresh_rebuilds_both_directories_and_keeps_borrowerless_selection(client, endpoint):
    sid = upload_guarantors(client)
    base = f"/api/sessions/{sid}"
    client.post(base + "/preflight")
    client.put(base + "/loan", json={"loan_id": "L2"})
    store = client.app.state.store
    session = store.get(sid)
    session.preflight.account_index_version = 1
    store.forget_account_candidates(sid)
    store.save(session)
    response = (client.put(base + "/loan", json={"loan_id": "L2"}) if endpoint == "loan"
                else client.get(base + "/" + endpoint))
    assert response.status_code == 200
    saved = store.get(sid)
    assert saved.preflight.account_index_version == ACCOUNT_INDEX_VERSION
    assert saved.selected_loan_id == "L2" and saved.selected_account_id is None
    assert store.loan_candidates(sid)[1] == 1
    assert store.account_candidates(sid)[1] == 0
    client.app.state.provider_stub.assert_not_called()


def test_directory_unlinked_missing_ambiguous_and_invalid_identity():
    records = [
        ("Account", {"Id": "A1", "Name": "First"}),
        ("Account", {"Id": "A2", "Name": "Second"}),
        ("Loan", {"Id": "UNLINKED"}),
        ("Loan", {"Id": "MISSING", "borrower_account_id": "ABSENT"}),
        ("Loan", {"Id": "VALID", "borrower_account_id": "A1"}),
        ("Loan", {"Id": "AMBIGUOUS"}),
        ("Loan", {"Id": "GUARANTOR"}),
        ("Loan", {"Id": "ROLE-MISSING"}),
        ("EntityInvolvement", {"Id": "J1", "loan_id": "AMBIGUOUS", "account_id": "A1", "role": "Borrower"}),
        ("EntityInvolvement", {"Id": "J2", "loan_id": "AMBIGUOUS", "account_id": "A2", "role": "Borrower"}),
        ("EntityInvolvement", {"Id": "ROOT-NOT-LOAN", "account_id": "A1", "role": "Guarantor"}),
        ("EntityInvolvement", {"Id": "J3", "loan_id": "ROLE-MISSING", "account_id": "A2"}),
        ("EntityInvolvement", {"Id": "J4", "loan_id": "GUARANTOR", "account_id": "A2", "role": "Guarantor"}),
        ("Loan", {"Id": False}), ("Loan", {"Id": 123}), ("Loan", {"Id": "=formula"}),
    ]
    doc = Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=0,
                   metadata={"records": [{"sheet": sheet, "row": index + 2, "values": values}
                                         for index, (sheet, values) in enumerate(records)]})
    loans = []
    count, accounts = _discover_accounts([doc], engine.load_underwriting_config(), loans)
    assert count == 6
    assert {item.loan_id: item.account_id for item in loans} == {
        "UNLINKED": None, "MISSING": None, "VALID": "A1", "AMBIGUOUS": None,
        "ROLE-MISSING": None, "GUARANTOR": None,
    }
    assert {item.account_id: item.loan_ids for item in accounts} == {
        "A1": ["AMBIGUOUS", "VALID"], "A2": ["AMBIGUOUS"],
    }
    graph = build_entity_graph([doc], engine.load_underwriting_config())
    assert graph.borrowers("GUARANTOR") == set()
    assert graph.borrowers("ROLE-MISSING") == set()
    for loan_id in ("GUARANTOR", "ROLE-MISSING", "AMBIGUOUS", "UNLINKED"):
        batch = engine.evaluate_underwriting([doc], selected_loan_id=loan_id)
        assert len(batch.loan_results) == 1
        assert batch.loan_results[0].borrower_account_id is None
        assert "explicit_unique_borrower_account_relationship" in batch.loan_results[0].missing_inputs


def test_store_atomic_replacement_optional_callers_prefix_search_and_cleanup(client):
    store = client.app.state.store
    accounts = [AccountCandidate(account_id="A1", name="Synthetic client", loan_ids=["L1", "L2"]),
                AccountCandidate(account_id="A2", name="Other", loan_ids=["L2"])]
    store.replace_account_candidates("synthetic", accounts)
    assert store.loan_candidate_account("synthetic", "L1") == "A1"
    assert store.loan_candidate_account("synthetic", "L2") is None
    assert store.loan_candidates("synthetic", "synthetic")[1] == 1
    assert store.loan_candidates("synthetic", "a1")[1] == 1
    assert store.loan_candidates("synthetic", "L")[1] == 2
    assert store.loan_candidates("synthetic", "%")[1] == 0
    assert store.loan_candidates("other-session")[1] == 0
    assert store.has_loan_candidate("synthetic", None, "L2")
    assert not store.has_loan_candidate("synthetic", "A1", "L2")
    with pytest.raises(sqlite3.IntegrityError):
        store.replace_account_candidates("synthetic", [], [LoanCandidate(loan_id="DUP"), LoanCandidate(loan_id="DUP")])
    assert store.loan_candidates("synthetic")[1] == 2
    assert store.account_candidates("synthetic")[1] == 2
    store.replace_account_candidates("synthetic", [], [LoanCandidate(loan_id="G_ONLY")])
    assert not store.loan_candidate_exists("synthetic", "L1")
    assert store.has_loan_candidate("synthetic", None, "G_ONLY")
    assert store.loan_candidates("synthetic", "G_")[1] == 1
    store.forget_account_candidates("synthetic")
    assert store.loan_candidates("synthetic")[1] == 0


def test_processing_rejects_nonexistent_loan_and_wrong_borrower(client):
    sid = upload_guarantors(client)
    base = f"/api/sessions/{sid}"
    client.put(base + "/loan", json={"loan_id": "L1"})
    store = client.app.state.store
    session = store.get(sid)
    session.selected_loan_id = "NONEXISTENT"
    store.save(session)
    assert "Enter a loan number" in client.post(base + "/process").json()["detail"]
    session.selected_loan_id = "L1"
    session.selected_account_id = "G1"
    store.save(session)
    assert client.post(base + "/process").status_code == 409
    session.selected_account_id = None
    session.preflight.account_index_version = 1
    store.save(session)
    assert client.post(base + "/process").status_code == 409
    client.app.state.provider_stub.assert_not_called()