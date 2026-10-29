"""Loan-first intake regressions using synthetic workbooks only."""

import io
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from backend.agents import DecisionAgent, ExtractionAgent, MemoryAgent
from backend.excel import extract_excel
from backend.llm import AnthropicLLM
from backend.loan_loading import load_loan_documents
from backend.main import create_app, process_session
from backend.models import Document, Report
from backend.tests.test_large_preflight import synthetic_documents, workbook_bytes


def canonical_workbook(sheets):
    book = Workbook(write_only=True)
    for name, headers, rows in sheets:
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_streaming_retains_only_requested_loan_and_late_linked_rows(tmp_path, monkeypatch):
    documents = synthetic_documents(2001)
    for document in documents:
        (tmp_path / document.id).write_bytes(workbook_bytes(document))
        document.metadata = {}
    import backend.loan_loading as loading
    parser = Mock(wraps=extract_excel)
    monkeypatch.setattr(loading, "extract_excel", parser)
    load_loan_documents(documents, tmp_path, "L02000")
    assert all(call.kwargs.get("record_filter") is not None for call in parser.call_args_list)
    for document in documents:
        assert len(document.metadata["records"]) == 1
        assert document.metadata["records"][0]["row"] == 2002
        assert "row-02000" in document.content
        assert "row-00000" not in document.content
        assert document.metadata["scoped_loan_id"] == "L02000"
    # Conflicting source attributes on the matching row must not be discarded.
    assert documents[2].metadata["records"][0]["values"]["fico"] == "620"


def test_header_validation_stops_before_reading_any_data_row(monkeypatch):
    raw = canonical_workbook([("Loan", ["Id"], [["L1"], ["L2"]])])
    original = ReadOnlyWorksheet.iter_rows

    def header_only(sheet, *args, **kwargs):
        rows = original(sheet, *args, **kwargs)
        yield next(rows)
        pytest.fail("Header validation must not consume data rows")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", header_only)
    _, metadata = extract_excel(raw, include_text=False, headers_only=True)
    assert metadata["headers_only"] is True
    assert metadata["records"] == []


def test_renamed_workbook_can_be_recognized_from_its_sheet_title(tmp_path):
    raw = canonical_workbook([("MTB Query 1", ["Loan_ID", "Account_ID", "Role"],
                              [["L1", "A1", "Borrower"], ["L2", "A2", "Borrower"]])])
    doc = Document(name="renamed.xlsx", extension=".xlsx", size=len(raw))
    (tmp_path / doc.id).write_bytes(raw)
    load_loan_documents([doc], tmp_path, "L1")
    assert len(doc.metadata["records"]) == 1
    assert doc.metadata["records"][0]["values"]["Loan_ID"] == "L1"


def test_relationship_closure_preserves_children_conflicts_and_excludes_other_loans(tmp_path):
    raw = canonical_workbook([
        ("Indicator", ["Id", "analysis_id", "value"], [["I1", "D1", "Y"], ["IX", "DX", "N"]]),
        ("DataSourceAnalysis", ["Id", "entity_id"], [["D1", "E1"], ["DX", "EX"]]),
        ("LCREntity", ["Id", "response_id", "account_id"], [["E1", "R1", "A1"], ["EX", "RX", "A1"]]),
        ("LCRESPONSE", ["Id", "loan_id"], [["R1", "L1"], ["RX", "LX"]]),
        ("CreditReport", ["Id", "account_id", "personal_fico"], [["C1", "A1", "700"], ["C1", "A1", "600"], ["CX", "AX", "800"]]),
        ("Account", ["Id", "name"], [["A1", "Related"], ["AX", "Unrelated"]]),
        ("Loan", ["Id", "borrower_account_id", "response_id"], [["L1", "A1", "R1"], ["L2", "A1", "R2"], ["LX", "AX", "RX"]]),
    ])
    doc = Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=len(raw))
    (tmp_path / doc.id).write_bytes(raw)
    load_loan_documents([doc], tmp_path, "L1")
    records = doc.metadata["records"]
    assert {row["values"]["Id"] for row in records} == {"I1", "D1", "E1", "R1", "C1", "A1", "L1"}
    assert len([row for row in records if row["sheet"] == "CreditReport"]) == 2
    assert "L2" not in doc.content and "Unrelated" not in doc.content


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(AnthropicLLM, "complete", Mock(side_effect=AssertionError("No provider during intake")))
    with TestClient(create_app(tmp_path / "data")) as active:
        yield active


def upload(client):
    sid = client.post("/api/sessions", json={}).json()["id"]
    documents = synthetic_documents(3)
    result = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", (doc.name, workbook_bytes(doc))) for doc in documents
    ])
    assert result.status_code == 200
    return sid, result.json()


def test_upload_asks_for_loan_and_preflight_reads_headers_only(client, monkeypatch):
    original = ExtractionAgent.extract

    def headers_only(self, document, path, **kwargs):
        assert kwargs["headers_only"] is True
        original(self, document, path, **kwargs)
        assert document.metadata["records"] == []

    monkeypatch.setattr(ExtractionAgent, "extract", headers_only)
    sid, session = upload(client)
    assert "What loan number" in session["messages"][-1]["content"]
    assert all(doc["metadata"] == {} for doc in session["documents"])
    result = client.post(f"/api/sessions/{sid}/preflight").json()["preflight"]
    assert result["requires_loan_number"] is True
    assert result["loan_count"] == 0
    assert client.post(f"/api/sessions/{sid}/process").status_code == 409


def test_loan_lookup_unknown_switching_and_scoped_processing(client, monkeypatch):
    sid, _ = upload(client)
    assert client.put(f"/api/sessions/{sid}/loan", json={"loan_id": "MISSING"}).status_code == 404
    assert client.app.state.store.get(sid).selected_loan_id is None
    selected = client.put(f"/api/sessions/{sid}/loan", json={"loan_id": "L00002"})
    assert selected.status_code == 200, selected.text
    assert selected.json()["preflight"]["loan_count"] == 1
    assert selected.json()["preflight"]["scoped_loan_id"] == "L00002"
    assert selected.json()["selected_account_id"] == "A00002"

    def decide(self, session, findings):
        assert all(len(doc.metadata["records"]) == 1 for doc in session.documents)
        assert all("row-00002" in doc.content and "row-00000" not in doc.content for doc in session.documents)
        return Report(decision="needs_review", summary="Synthetic review", findings=[], document_count=3)

    monkeypatch.setattr(ExtractionAgent, "extract", Mock(side_effect=AssertionError("No full extraction")))
    monkeypatch.setattr(DecisionAgent, "decide", decide)
    monkeypatch.setattr(MemoryAgent, "index", lambda self, session: "Synthetic index")
    process_session(client.app.state.store, sid)
    assert client.app.state.store.get(sid).status == "completed"
    monkeypatch.undo()
    switched = client.put(f"/api/sessions/{sid}/loan", json={"loan_id": "L00000"})
    assert switched.status_code == 200, switched.text
    assert switched.json()["selected_loan_id"] == "L00000"
    assert switched.json()["report"] is None
    assert all(doc["metadata"] == {} for doc in switched.json()["documents"])


def test_chat_requests_loan_number_before_loading_or_calling_provider(client):
    sid, _ = upload(client)
    prompt = client.post(f"/api/sessions/{sid}/messages", json={"content": "Please review these files"})
    assert prompt.status_code == 200
    assert "What loan number" in prompt.json()["messages"][-1]["content"]
    assert all(doc["metadata"] == {} for doc in prompt.json()["documents"])
    selected = client.post(f"/api/sessions/{sid}/messages", json={"content": "loan number L00002"})
    assert selected.status_code == 200, selected.text
    assert selected.json()["selected_loan_id"] == "L00002"
    assert selected.json()["messages"][-2]["role"] == "user"
    assert "Selected explicit loan number L00002" in selected.json()["messages"][-1]["content"]
