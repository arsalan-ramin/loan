"""Renamed MTB data must never silently fall through to generic approval."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.agents import DecisionAgent, RulesAgent
from backend.datasets import (SIGNATURES, SOQL_ROOT_SIGNATURES, account_reference_headers,
                              borrower_role_headers, loan_reference_headers,
                              mtb_query_ids, soql_account_reference_headers)
from backend.llm import AnthropicLLM, LLMError
from backend.main import create_app
from backend.models import Document, Report, Session
from backend.policy import evaluate_session
from backend.storage import Store


@pytest.mark.parametrize("name,expected", [
    ("MTB_Query_1.xlsx", {1}), ("MTB_Query_1_Sample_Data.xlsx", {1}),
    ("MTB_Query_2_Sample_Data.xlsx", {2}), ("MTB_Query_3_Sample_Data.xlsx", {3}),
    ("MTB_Query_4_Sample_Data.xlsx", {4}),
    ("mtb query 4 (copy).XLSX", {4}), ("MTB-Query-2-test.xlsx", {2}),
    ("MTB_Query_14.xlsx", set()),
    ("MTB_Query_1.txt", set()), ("MTB_Query_1.xlsm", set()), ("report.xlsx", set()),
])
def test_sample_names_and_unrelated_files(name, expected):
    assert mtb_query_ids(name, {}) == expected


@pytest.mark.parametrize("query", [1, 2, 4])
def test_renamed_workbook_detected_from_sheet_or_headers(query):
    assert mtb_query_ids("renamed.xlsx", {"sheets": [{"name": f"MTB Query {query}", "headers": []}]}) == {query}
    assert mtb_query_ids("renamed.xlsx", {"sheets": [{"name": "Data", "headers": list(SIGNATURES[query])}]}) == {query}
    assert mtb_query_ids("renamed.xlsx", {"sheets": [{"name": "README", "is_reference": True,
                                                         "headers": list(SIGNATURES[query])}]}) == set()


def test_query_one_bureau_columns_not_misclassified_as_query_four():
    metadata = {"sheets": [{"name": "Data", "headers": list(SIGNATURES[1] | SIGNATURES[4])}]}
    assert mtb_query_ids("renamed.xlsx", metadata) == {1}


@pytest.mark.parametrize("query", [1, 2, 4])
def test_copied_soql_headers_recognize_renamed_workbooks(query):
    headers = set(SOQL_ROOT_SIGNATURES[query][1]) | {"Id"}
    if query == 2:
        headers.add("SCOC__LCRDataSourceAnalysis__r.SCOC__LCREntity__r."
                    "SCOC__Entity_Involvement__r.LLC_BI__Loan__c")
    metadata = {"sheets": [{"name": "Sheet1", "headers": sorted(headers)}]}
    assert mtb_query_ids("export.xlsx", metadata) == {query}
    metadata["sheets"][0]["is_reference"] = True
    assert mtb_query_ids("export.xlsx", metadata) == set()


def test_nested_query_one_headers_take_precedence_over_bureau_signature():
    headers = SOQL_ROOT_SIGNATURES[1][1] | SIGNATURES[4]
    assert mtb_query_ids("export.xlsx", {"sheets": [{"name": "Sheet1", "headers": list(headers)}]}) == {1}


def test_partial_copied_headers_do_not_guess_query_role():
    headers = ["Id", "SCOC__IndicatorKey__c", "SCOC__IndicatorValue__c"]
    assert mtb_query_ids("export.xlsx", {"sheets": [{"name": "Sheet1", "headers": headers}]}) == set()


def test_canonical_schema_coverage_uses_manifest_before_soql_signatures():
    from backend.datasets import soql_mapping_coverage
    from backend.underwriting_engine import load_underwriting_config
    manifest = load_underwriting_config()["mapping_manifest"]
    headers = ["Id", "group_credit_usd", "real_estate_secured", "requested_path", "unsupported_attribute"]
    mapped, targets = soql_mapping_coverage(1, headers, manifest, sheet_name="Loan")
    assert mapped == set(headers) - {"unsupported_attribute"}
    assert {"Loan.id", "Loan.group_credit_usd", "Loan.real_estate_secured", "Loan.requested_path"} <= targets
    mapped, targets = soql_mapping_coverage(4, ["Id", "bankruptcy_count", "report_date"], manifest, sheet_name="CreditReport")
    assert "bankruptcy_count" in mapped and "CreditReport.report_date" in targets


@pytest.mark.parametrize("header", [
    "ACCOUNT_ID", "Account ID", "Account-ID", "Account Column", "Account ID Column",
    "LLC_BI__Account__c", "LLC_BI__Account__r.Id", "nFUSE__Account__c",
])
def test_account_reference_header_variants_are_explicitly_normalized(header):
    assert account_reference_headers([header]) == [header]


def test_account_names_and_generic_ids_are_not_account_relationship_keys():
    assert account_reference_headers(["Name", "Account Name", "Id", "Record_ID"]) == []


def test_namespaced_soql_loan_and_borrower_relationship_headers_are_explicit():
    prefix = "SCOC__LCRDataSourceAnalysis__r.SCOC__LCREntity__r.SCOC__Entity_Involvement__r."
    loan_header = prefix + "LLC_BI__Loan__c"
    role_header = prefix + "LLC_BI__Borrower_Type__c"
    assert loan_reference_headers([loan_header, "Unrelated_Loan_Description__c", "Id"]) == [loan_header]
    assert borrower_role_headers([role_header, "Borrower_Name__c", "Name"]) == [role_header]


def test_soql_account_foreign_keys_follow_query_specific_relationship_contract():
    prefix = "SCOC__LCRDataSourceAnalysis__r.SCOC__LCREntity__r.SCOC__Entity_Involvement__r."
    involvement_account = prefix + "LLC_BI__Account__c"
    q4_account = "nFUSE__Account__c"
    distractors = ["LLC_BI__Account__c", "LLC_BI__Account__r.Id", "ACCOUNT_ID"]

    assert soql_account_reference_headers(1, [involvement_account, *distractors]) == [involvement_account]
    assert soql_account_reference_headers(2, [involvement_account, *distractors]) == [involvement_account]
    assert soql_account_reference_headers(4, [q4_account, *distractors]) == [q4_account]
    assert soql_account_reference_headers(1, [q4_account]) == []
    assert soql_account_reference_headers(4, [involvement_account]) == []


def test_sample_report_runs_guideline_not_word_count(monkeypatch):
    monkeypatch.setattr(AnthropicLLM, "complete", lambda *_: (_ for _ in ()).throw(LLMError("Network disabled")))
    documents = [Document(name=f"MTB_Query_{query}_Sample_Data.xlsx", extension=".xlsx", size=10,
                          content="Many readable words. " * 50,
                          metadata={"data_row_count": 1, "records": [{"sheet": f"MTB Query {query}", "row": 2,
                          "values": {"nFUSE__Credit_Report_Freeze__c": True},
                          "sources": {"nFUSE__Credit_Report_Freeze__c": "A2"}}]}) for query in (1, 2, 4)]
    session = Session(documents=documents)
    findings = RulesAgent().evaluate(session)
    report = DecisionAgent().decide(session, findings)
    assert report.decision == "needs_review"
    assert report.policy_evaluation is not None
    # V4 requires explicit loan identity: unrelated source rows are not profiles.
    assert report.underwriting is not None
    assert report.underwriting.loan_results == []
    assert report.policy_evaluation.profiles == []
    assert report.policy_evaluation.rule_results == []
    assert "NO_EXPLICIT_LOAN_RECORDS" in report.underwriting.data_quality_issues
    assert report.policy_evaluation.source_document == report.underwriting.source_document
    assert not any(f.rule_id in {"readable-content", "draft-check"} for f in findings)
    assert not any("Required workbook" in item for item in report.policy_evaluation.missing_information)


def test_stale_sample_approval_invalidated_even_without_default_checks(tmp_path):
    root = tmp_path / "data"
    store = Store(root)
    session = Session(documents=[Document(name="MTB_Query_1_Sample_Data.xlsx", extension=".xlsx", size=10)], rules=[], status="completed",
                      report=Report(decision="approved", summary="Generic approval", findings=[], document_count=1,
                                    llm_analysis="Generic checks passed."))
    store.save(session)
    with TestClient(create_app(root)) as client:
        migrated = client.get(f"/api/sessions/{session.id}").json()
        assert migrated["status"] == "ready"
        assert migrated["report"] is None
        assert migrated["policy_rules"]
        assert "invalidated" in migrated["messages"][-1]["content"]
        assert client.get(f"/api/sessions/{session.id}/report").status_code == 409


def test_duplicate_query_files_disclosed_not_joined():
    session = Session(documents=[Document(name=name, extension=".xlsx", size=10) for name in
                                 ("MTB_Query_1.xlsx", "MTB_Query_1_Sample_Data.xlsx")])
    policy = evaluate_session(session)
    assert policy is not None
    assert any("Multiple workbooks" in item for item in policy.missing_information)
    assert any("MTB Query 2" in item for item in policy.missing_information)


def test_header_recognition_validates_structure_without_loading_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(AnthropicLLM, "complete", lambda *_: (_ for _ in ()).throw(LLMError("Network disabled")))
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(sorted(SIGNATURES[4]))
    sheet.append([False, 0, False])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    app = create_app(tmp_path / "data")
    with TestClient(app) as client:
        sid = client.post("/api/sessions", json={}).json()["id"]
        client.post(f"/api/sessions/{sid}/documents", files=[("files", ("renamed.xlsx", stream.getvalue()))])
        result = client.post(f"/api/sessions/{sid}/preflight").json()
        assert result["status"] == "ready"
        assert result["preflight"]["recognized_mtb"] is True
        assert result["preflight"]["query_roles"] == [4]
        assert result["preflight"]["requires_loan_number"] is True
        assert all(doc["metadata"] == {} for doc in result["documents"])
        assert result["report"] is None
        assert client.post(f"/api/sessions/{sid}/process").status_code == 409