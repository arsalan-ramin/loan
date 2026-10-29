"""Contract and workflow tests against real temporary SQLite and Chroma stores."""

import ast
import io
import inspect
import json
import zipfile

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf._page import PageObject

from backend import agents, main
from backend.agents import ExtractionAgent
from backend.main import create_app, process_session
from backend.models import Document
from backend.storage import embed
from backend.llm import AnthropicLLM


@pytest.fixture(autouse=True)
def mock_provider(monkeypatch):
    """Workflow tests exercise retrieval/persistence without making billable API calls."""
    def complete(_self, _system, prompt):
        if prompt.startswith("SESSION DATA"):
            context = json.loads(prompt.split("\n", 1)[1].split("\n\nCURRENT USER QUESTION:")[0])
            return "\n".join(
                [f'[{name}] {text}' for name, text in context["document_excerpts"]] +
                [m["content"] for m in context["recent_messages"] if m["role"] == "user"]
            ) or "Please upload your required files first."
        return "Claude commentary based on the supplied rule findings."
    monkeypatch.setattr(AnthropicLLM, "complete", complete)


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / "data")) as client:
        yield client


def new_session(client):
    response = client.post("/api/sessions", json={})
    assert response.status_code == 201
    return response.json()["id"]


def upload(client, session_id, name="evidence.txt", content=b"Approved policy. " * 40):
    return client.post(f"/api/sessions/{session_id}/documents", files=[("files", (name, content))])


def finish(client, session_id):
    # Run deterministically in tests, while exercising the actual worker implementation.
    store = client.app.state.store
    session = store.get(session_id)
    session.status = "processing"
    store.save(session)
    process_session(store, session_id)
    return client.get(f"/api/sessions/{session_id}").json()


def test_upload_first_and_health(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert client.get("/openapi.json").json()["info"]["title"] == "Agent document intelligence"
    sid = new_session(client)
    session = client.get(f"/api/sessions/{sid}").json()
    assert "upload" in session["messages"][0]["content"].lower()
    assert client.post(f"/api/sessions/{sid}/process").status_code == 400
    assert client.get(f"/api/sessions/{sid}/report").status_code == 409


def test_successful_review_and_html(client):
    sid = new_session(client)
    assert upload(client, sid).status_code == 200
    session = finish(client, sid)
    assert session["status"] == "completed", session["error"]
    assert session["report"]["decision"] == "approved"
    assert session["documents"][0]["metadata"]["word_count"] == 80
    assert all(step["status"] == "completed" for step in session["steps"])
    report = client.get(f"/api/sessions/{sid}/report")
    assert report.headers["content-disposition"] == f'attachment; filename="agent-report-{sid}.html"'
    assert "<title>Agent · " in report.text
    assert "AGENT / DOCUMENT INTELLIGENCE" in report.text
    assert "Generated locally by Agent." in report.text
    assert "Decision evidence" in report.text


def test_all_rule_types_and_rejection(client):
    sid = new_session(client)
    upload(client, sid, content=b"DRAFT policy")
    rules = [
        {"id": "1", "name": "Required", "kind": "required_text", "value": "approved", "severity": "fail"},
        {"id": "2", "name": "No draft", "kind": "forbidden_text", "value": "draft", "severity": "review"},
        {"id": "3", "name": "Length", "kind": "min_words", "value": "30", "severity": "review"},
        {"id": "4", "name": "Author", "kind": "required_metadata", "value": "author", "severity": "review"},
        {"id": "5", "name": "Disabled", "kind": "required_text", "value": "missing", "enabled": False},
    ]
    assert client.put(f"/api/sessions/{sid}/rules", json={"rules": rules}).status_code == 200
    session = finish(client, sid)
    assert session["report"]["decision"] == "rejected"
    assert len(session["report"]["findings"]) == 4
    assert "DRAFT" in session["report"]["findings"][1]["evidence"]


def test_review_decision(client):
    sid = new_session(client)
    upload(client, sid, content=b"short document")
    assert finish(client, sid)["report"]["decision"] == "needs_review"


@pytest.mark.parametrize("name,content", [("bad.exe", b"binary"), ("empty.txt", b"")])
def test_invalid_upload(client, name, content):
    sid = new_session(client)
    assert upload(client, sid, name, content).status_code == 400
    assert client.get(f"/api/sessions/{sid}").json()["documents"] == []


def test_invalid_rule_and_empty_chat(client):
    sid = new_session(client)
    rule = {"id": "bad", "name": "Length", "kind": "min_words", "value": "NaN"}
    assert client.put(f"/api/sessions/{sid}/rules", json={"rules": [rule]}).status_code == 422
    assert client.post(f"/api/sessions/{sid}/messages", json={"content": "   "}).status_code == 422


def test_conversation_memory_and_isolation(client):
    first, second = new_session(client), new_session(client)
    client.post(f"/api/sessions/{first}/messages", json={"content": "My project code is ORCHID-ALPHA."})
    answer = client.post(f"/api/sessions/{first}/messages", json={"content": "Remember the project code I said earlier?"}).json()
    assert "ORCHID-ALPHA" in answer["messages"][-1]["content"]
    other = client.post(f"/api/sessions/{second}/messages", json={"content": "Remember ORCHID-ALPHA from earlier?"}).json()
    assert "My project code is" not in other["messages"][-1]["content"]


def test_document_retrieval_and_invalidation(client):
    sid = new_session(client)
    upload(client, sid, content=b"ORCHID policy evidence. " * 40)
    session = finish(client, sid)
    answer = client.post(f"/api/sessions/{sid}/messages", json={"content": "Find ORCHID policy evidence"}).json()
    assert "[evidence.txt]" in answer["messages"][-1]["content"]
    doc_id = session["documents"][0]["id"]
    removed = client.delete(f"/api/sessions/{sid}/documents/{doc_id}").json()
    assert removed["status"] == "awaiting_upload"
    assert removed["report"] is None
    assert client.app.state.store.recall(sid, "ORCHID", "document") == []


def test_processing_locks_edits(client):
    sid = new_session(client)
    upload(client, sid)
    store = client.app.state.store
    session = store.get(sid)
    session.status = "processing"
    store.save(session)
    assert upload(client, sid).status_code == 409
    assert client.put(f"/api/sessions/{sid}/rules", json={"rules": []}).status_code == 409
    assert client.post(f"/api/sessions/{sid}/process").status_code == 409


def test_report_escapes_uploaded_text(client):
    sid = new_session(client)
    upload(client, sid, content=b"<script>alert('bad')</script> DRAFT")
    finish(client, sid)
    html = client.get(f"/api/sessions/{sid}/report").text
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


def test_malformed_json_failure_is_retryable(client):
    sid = new_session(client)
    upload(client, sid, "bad.json", b"{not json}")
    session = finish(client, sid)
    assert session["status"] == "failed"
    assert session["report"] is None
    assert any(s["status"] == "failed" for s in session["steps"])


def test_csv_json_docx_extraction(tmp_path):
    extractor = ExtractionAgent()
    for suffix, raw, expected in [
        (".csv", b"name,amount\nAlice,30\nBob,40", "rows"),
        (".json", b'{"project":"Orchid"}', "json_keys"),
    ]:
        path = tmp_path / f"sample{suffix}"
        path.write_bytes(raw)
        doc = Document(name=path.name, size=len(raw), extension=suffix)
        extractor.extract(doc, path)
        assert expected in doc.metadata
    word = WordDocument()
    word.add_paragraph("This is a real DOCX sample.")
    word.core_properties.author = "Agent Test"
    stream = io.BytesIO()
    word.save(stream)
    path = tmp_path / "sample.docx"
    path.write_bytes(stream.getvalue())
    doc = Document(name=path.name, size=path.stat().st_size, extension=".docx")
    extractor.extract(doc, path)
    assert "real DOCX" in doc.content
    assert doc.metadata["author"] == "Agent Test"


def test_persistence_and_restart_recovery(tmp_path):
    root = tmp_path / "persistent"
    app = create_app(root)
    with TestClient(app) as client:
        sid = new_session(client)
        client.post(f"/api/sessions/{sid}/messages", json={"content": "Remember SUNFLOWER project."})
        session = app.state.store.get(sid)
        session.status = "processing"
        app.state.store.save(session)
    restarted_app = create_app(root)
    with TestClient(restarted_app) as client:
        session = client.get(f"/api/sessions/{sid}").json()
        assert session["status"] == "failed"
        assert "restarted" in session["error"]
        assert any("SUNFLOWER" in m["content"] for m in session["messages"])
        assert restarted_app.state.store.recall(sid, "SUNFLOWER", "conversation")


def test_actual_async_process_endpoint(client):
    sid = new_session(client)
    upload(client, sid)
    response = client.post(f"/api/sessions/{sid}/process")
    assert response.status_code == 202
    assert response.json()["status"] == "processing"
    # Wait for submitted work without polling or sleeping.
    client.app.state.executor.shutdown(wait=True)
    assert client.get(f"/api/sessions/{sid}").json()["status"] == "completed"


def test_embeddings_are_stable_and_normalized():
    vector = embed("orchid policy orchid")
    assert vector == embed("ORCHID policy orchid")
    assert len(vector) == 384
    assert abs(sum(x * x for x in vector) - 1) < 1e-9


def test_pdf_extraction_and_encryption(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata({"/Author": "Agent PDF Test", "/Title": "Scanned sample"})
    path = tmp_path / "sample.pdf"
    writer.write(path)
    doc = Document(name=path.name, size=path.stat().st_size, extension=".pdf")
    ExtractionAgent().extract(doc, path)
    assert doc.metadata["pages"] == 1
    assert doc.metadata["author"] == "Agent PDF Test"
    assert "OCR" in doc.metadata["warning"]
    writer.encrypt("test-password")
    writer.write(path)
    with pytest.raises(ValueError, match="encrypted PDFs"):
        ExtractionAgent().extract(doc, path)


@pytest.mark.parametrize("name,raw", [("empty.txt", b""), ("unsupported.exe", b"invalid")])
def test_invalid_batch_rolls_back_only_new_files(client, name, raw):
    sid = new_session(client)
    assert upload(client, sid, content=b"existing source").status_code == 200
    before = client.get(f"/api/sessions/{sid}").json()
    directory = client.app.state.store.root / "uploads" / sid
    existing = {path.name: path.read_bytes() for path in directory.iterdir()}
    response = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", ("small.txt", b"valid file")),
        ("files", (name, raw)),
    ])
    assert response.status_code == 400
    assert client.get(f"/api/sessions/{sid}").json() == before
    assert {path.name: path.read_bytes() for path in directory.iterdir()} == existing


@pytest.mark.parametrize("extension", [".txt", ".xlsx"])
def test_upload_ignores_retired_size_environment_settings(client, monkeypatch, extension):
    monkeypatch.setenv("FOLIO_MAX_FILE_MB", "1")
    monkeypatch.setenv("FOLIO_MAX_XLSX_MB", "1")
    sid = new_session(client)
    raw = b"x" * (1024 * 1024 + 1)
    response = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", (f"large{extension}", raw)),
    ])
    assert response.status_code == 200
    document, = response.json()["documents"]
    assert document["size"] == len(raw)
    assert (client.app.state.store.root / "uploads" / sid / document["id"]).read_bytes() == raw


def test_generic_upload_has_no_conversation_or_multipart_file_count_cap(client):
    sid = new_session(client)
    assert upload(client, sid).status_code == 200
    response = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", (f"source-{index}.txt", b"synthetic")) for index in range(1001)
    ])
    assert response.status_code == 200
    assert len(response.json()["documents"]) == 1002


def test_upload_still_requires_files_and_documents_openapi(client):
    sid = new_session(client)
    assert client.post(f"/api/sessions/{sid}/documents").status_code == 400
    schema = client.get("/openapi.json").json()["paths"]["/api/sessions/{session_id}/documents"]["post"]
    files = schema["requestBody"]["content"]["multipart/form-data"]["schema"]["properties"]["files"]
    assert files == {"type": "array", "items": {"type": "string", "format": "binary"}}


@pytest.mark.parametrize("first,second", [(".txt", ".xlsx"), (".xlsx", ".txt")])
def test_upload_still_rejects_mixed_workbook_and_generic_sources(client, first, second):
    sid = new_session(client)
    assert upload(client, sid, name=f"first{first}").status_code == 200
    before = client.get(f"/api/sessions/{sid}").json()
    assert upload(client, sid, name=f"second{second}").status_code == 400
    assert client.get(f"/api/sessions/{sid}").json() == before
    other = new_session(client)
    response = client.post(f"/api/sessions/{other}/documents", files=[
        ("files", ("source.xlsx", b"synthetic")), ("files", ("source.txt", b"synthetic")),
    ])
    assert response.status_code == 400
    assert client.get(f"/api/sessions/{other}").json()["documents"] == []


def test_file_resource_guard_code_is_removed():
    agent_tree = ast.parse(inspect.getsource(agents))
    assert not any(isinstance(node, ast.Name) and node.id in {"MAX_CONTENT", "zipfile"}
                   for node in ast.walk(agent_tree))
    main_tree = ast.parse(inspect.getsource(main))
    assert not any(isinstance(node, ast.Constant) and node.value in
                   ("FOLIO_MAX_FILE_MB", "FOLIO_MAX_XLSX_MB") for node in ast.walk(main_tree))


@pytest.mark.parametrize("extension", [".txt", ".md", ".csv", ".json", ".docx"])
def test_generic_extraction_accepts_more_than_two_million_characters(tmp_path, extension):
    text = "synthetic " * 200_001
    if extension == ".docx":
        word = WordDocument()
        word.add_paragraph(text)
        stream = io.BytesIO()
        word.save(stream)
        raw = stream.getvalue()
    elif extension == ".json":
        raw = json.dumps({"description": text}).encode()
    else:
        raw = text.encode()
    path = tmp_path / f"synthetic{extension}"
    path.write_bytes(raw)
    document = Document(name=path.name, extension=extension, size=len(raw))
    ExtractionAgent().extract(document, path)
    assert document.metadata["character_count"] == len(document.content) > 2_000_000
    assert text in document.content


def test_pdf_accepts_more_than_500_pages_and_two_million_characters(tmp_path, monkeypatch):
    writer = PdfWriter()
    for _ in range(501):
        writer.add_blank_page(width=72, height=72)
    path = tmp_path / "synthetic.pdf"
    writer.write(path)
    # Real 501-page PDF, synthetic page text: no large PDF asset or OCR/provider call.
    monkeypatch.setattr(PageObject, "extract_text", lambda self: "synthetic " * 400)
    document = Document(name=path.name, extension=".pdf", size=path.stat().st_size)
    ExtractionAgent().extract(document, path)
    assert document.metadata["pages"] == 501
    assert document.metadata["character_count"] == len(document.content) > 2_000_000


def test_invalid_docx_package_is_still_rejected(tmp_path):
    path = tmp_path / "invalid.docx"
    path.write_bytes(b"not a ZIP package")
    document = Document(name=path.name, extension=".docx", size=path.stat().st_size)
    with pytest.raises(zipfile.BadZipFile):
        ExtractionAgent().extract(document, path)