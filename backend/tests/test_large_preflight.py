"""Synthetic-only large workbook regression; never reads workspace uploads."""

import io
from time import perf_counter
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import backend.assistant_workflow as workflow
from backend.agents import ExtractionAgent
from backend.llm import AnthropicLLM
from backend.main import create_app
from backend.models import Document
from backend.underwriting_engine import load_underwriting_config


ROW_COUNT = 16_001
JUNCTION = "SCOC__LCRDataSourceAnalysis__r.SCOC__LCREntity__r.SCOC__Entity_Involvement__r."


def synthetic_documents(count=ROW_COUNT):
    """Distinct explicit loans/borrowers, nested SOQL columns, and raw extras."""
    documents = []
    for query in (1, 2, 4):
        records = []
        for index in range(count):
            account, loan = f"A{index:05d}", f"L{index:05d}"
            if query in (1, 2):
                values = {
                    "Id": f"Q{query}-{index:05d}",
                    JUNCTION + "LLC_BI__Loan__c": loan,
                    JUNCTION + "LLC_BI__Account__c": account,
                    JUNCTION + "LLC_BI__Borrower_Type__c": "Borrower",
                    JUNCTION + "LLC_BI__Account__r.Name": f"Synthetic client {index:05d}",
                    JUNCTION + "LLC_BI__Loan__r.LLC_BI__Amount__c": "100000.01",
                }
                if query == 1:
                    values.update({"SCOC__CharacteristicKey__c": "numCollection",
                                   "SCOC__CharacteristicValue__c": "0"})
                else:
                    values.update({"SCOC__IndicatorKey__c": "thinFileIndicator",
                                   "SCOC__IndicatorValue__c": "N"})
            else:
                values = {"Id": f"R{index:05d}", "nFUSE__Account__c": account,
                          "nFUSE__Credit_Score__c": "740", "nFUSE__Status__c": "Completed",
                          # Full graph must still detect a defect on the final row.
                          "fico": "620" if index == count - 1 else None,
                          "nFUSE__Credit_Report_Freeze__c": "N",
                          "nFUSE__OFAC_Alert_Indicator__c": "N",
                          "nFUSE__Number_Of_Bankruptcies__c": "0"}
            values.update({"Unsupported_Source_Field__c": f"raw-query-{query}-row-{index:05d}",
                           "Optional_Blank__c": None})
            records.append({"sheet": "Sheet1", "row": index + 2, "values": values})
        documents.append(Document(
            name=f"MTB_Query_{query}_SOQL_Only.xlsx", extension=".xlsx", size=0,
            metadata={"sheets": [{"name": "Sheet1", "headers": list(records[0]["values"]),
                                    "data_rows": count}], "records": records},
        ))
    return documents


def workbook_bytes(document):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Sheet1")
    headers = document.metadata["sheets"][0]["headers"]
    sheet.append(headers)
    for record in document.metadata["records"]:
        sheet.append([record["values"].get(header) for header in headers])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_discovery_compiles_headers_per_schema_not_per_row(monkeypatch):
    documents = synthetic_documents(200)
    plan = Mock(wraps=workflow._mapping_plan)
    root = Mock(wraps=workflow.soql_root_table)
    suffixes = Mock(wraps=workflow.soql_field_headers)
    monkeypatch.setattr(workflow, "_mapping_plan", plan)
    monkeypatch.setattr(workflow, "soql_root_table", root)
    monkeypatch.setattr(workflow, "soql_field_headers", suffixes)
    row_mapper = Mock(side_effect=AssertionError("Do not rebuild mappings per row"))
    monkeypatch.setattr(workflow, "_mapped_values", row_mapper)
    count, candidates = workflow._discover_accounts(documents, load_underwriting_config())
    assert count == len(candidates) == 200
    assert candidates[-1].loan_ids == ["L00199"]
    assert plan.call_count == 9  # 3 tables x 3 workbook schemas, not 1800.
    assert root.call_count == 3
    assert suffixes.call_count < 150
    row_mapper.assert_not_called()


@pytest.mark.parametrize("metadata_headers", [True, False])
def test_discovery_sparse_rows_conflicts_roles_names_and_invalid_ids(metadata_headers):
    config = load_underwriting_config()
    base = {
        "Id": "J1", JUNCTION + "LLC_BI__Loan__c": "L1",
        JUNCTION + "LLC_BI__Account__c": "A1",
        JUNCTION + "LLC_BI__Borrower_Type__c": "Borrower",
        JUNCTION + "LLC_BI__Account__r.Name": "  Synthetic name  ",
        "SCOC__CharacteristicKey__c": "numCollection", "SCOC__CharacteristicValue__c": "0",
    }
    rows = [base,
            {**base, JUNCTION + "LLC_BI__Account__r.Name": ""},
            {**base, JUNCTION + "LLC_BI__Account__r.Name": None},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L2", "Loan_ID": "OTHER"},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L3",
             JUNCTION + "LLC_BI__Borrower_Type__c": "Guarantor"},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L4",
             JUNCTION + "LLC_BI__Account__c": 123},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L5", "Loan_ID": ""},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L6", "Loan_ID": None},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L7", "Loan_ID": "L7"},
            # Conflicting Account/role aliases suppress mapped identity/role.
            {**base, JUNCTION + "LLC_BI__Loan__c": "L8",
             "Other.Entity_Involvement__r.LLC_BI__Account__c": "A-OTHER"},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L9", "role": "Guarantor"},
            {**base, JUNCTION + "LLC_BI__Loan__c": "L10",
             JUNCTION + "LLC_BI__Borrower_Type__c": " "},
            # Header order changes and omitted blank keys are distinct schemas.
            dict(reversed(list(base.items())))]
    doc = Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=0,
                   metadata={"records": [{"sheet": "Sheet1", "values": values} for values in rows]})
    if metadata_headers:
        doc.metadata["sheets"] = [{"name": "Sheet1", "headers": list(dict.fromkeys(
            key for row in rows for key in row))}]
    count, candidates = workflow._discover_accounts([doc], config)
    # Existing discovery semantics: an involvement's own loan_id is resolved
    # independently of Loan.id; a conflicting Loan_ID does not erase that link.
    assert count == 10  # All explicit loans, even invalid/missing borrower identity or role.
    assert [candidate.model_dump() for candidate in candidates] == [{
        "account_id": "A1", "name": "Synthetic name", "loan_count": 5,
        "loan_ids": ["L1", "L2", "L5", "L6", "L7"],
    }]


def test_discovery_canonical_direct_borrowers_and_non_soql_fallback():
    # Preserve the existing canonical-sheet mapper's behavior, including direct
    # borrower FKs without a role, scalar ID checks, and later nonblank names.
    config = load_underwriting_config()
    records = [
        {"sheet": "Account", "values": {"Id": "A1", "name": "  First  "}},
        {"sheet": "Account", "values": {"Id": "A1", "name": ""}},
        {"sheet": "Loan", "values": {"Id": "L1", "borrower_account_id": "A1"}},
        {"sheet": "Loan", "values": {"Id": "L2", "borrower_account_id": None}},
        {"sheet": "Loan", "values": {"Id": "L3", "borrower_account_id": False}},
        {"sheet": "Loan", "values": {"Id": "L4", "Loan_ID": "OTHER", "borrower_account_id": "A1"}},
        {"sheet": "Account", "values": {"Id": "A1", "name": "  Last  "}},
        None, {"sheet": None, "values": {}}, {"sheet": "Loan", "values": []},
    ]
    doc = Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=0, metadata={"records": records})
    count, candidates = workflow._discover_accounts([doc], config)
    assert count == 3  # L2/L3 have no valid borrower, but are still explicit loans.
    assert [(item.account_id, item.name, item.loan_ids) for item in candidates] == [("A1", "Last", ["L1"])]


def test_large_three_workbook_loan_first_loading_preserves_sources(tmp_path, monkeypatch):
    provider = Mock(side_effect=AssertionError("Preflight/search/select must stay local"))
    monkeypatch.setattr(AnthropicLLM, "complete", provider)
    timings = {"extract": 0.0, "directory": 0.0, "graph": 0.0}

    def timed(name, function):
        def run(*args, **kwargs):
            started = perf_counter()
            try:
                return function(*args, **kwargs)
            finally:
                timings[name] += perf_counter() - started
        return run

    extract = ExtractionAgent.extract
    monkeypatch.setattr(ExtractionAgent, "extract", timed("extract", extract))
    monkeypatch.setattr(workflow, "_discover_accounts", timed("directory", workflow._discover_accounts))
    # Graph validation must receive only the requested loan's related records.
    graph_builder = Mock(wraps=timed("graph", workflow.build_entity_graph))
    monkeypatch.setattr(workflow, "build_entity_graph", graph_builder)
    documents = synthetic_documents()
    payloads = [workbook_bytes(document) for document in documents]
    del documents
    with TestClient(create_app(tmp_path / "data")) as client:
        submit = Mock(side_effect=AssertionError("Local endpoints must not start processing"))
        monkeypatch.setattr(client.app.state.executor, "submit", submit)
        sid = client.post("/api/sessions", json={}).json()["id"]
        uploaded = client.post(f"/api/sessions/{sid}/documents", files=[
            ("files", (f"MTB_Query_{query}_SOQL_Only.xlsx", payload))
            for query, payload in zip((1, 2, 4), payloads)
        ])
        assert uploaded.status_code == 200
        original_documents = uploaded.json()["documents"]
        started = perf_counter()
        response = client.post(f"/api/sessions/{sid}/preflight")
        timings["preflight_api"] = perf_counter() - started
        print("\nPreflight phase seconds:", {key: round(value, 3) for key, value in timings.items()})
        assert response.status_code == 200
        session = response.json()
        result = session["preflight"]
        assert result["query_roles"] == [1, 2, 4]
        assert result["query_file_set_valid"] is True
        assert result["loan_count"] == result["account_count"] == 0
        assert result["requires_loan_number"] is True
        assert result["external_processing"] is False
        assert "account_candidates" not in result
        assert "loan_candidates" not in result
        assert len(result["priority_issues"]) <= 10
        # Session includes a fixed policy-rule catalog (~60KB), but no raw rows.
        assert len(response.content) < 80_000
        assert session["report"] is None and session["status"] == "ready"
        assert all("records" not in doc["metadata"] and not doc.get("content")
                   for doc in session["documents"])
        assert graph_builder.call_count == 0

        started = perf_counter()
        page = client.get(f"/api/sessions/{sid}/loans")
        assert page.status_code == 200
        assert page.json()["total"] == 0
        assert page.json()["results"] == []
        assert len(page.content) < 10_000
        # Exact lookup reaches beyond 15k without a portfolio-wide directory.
        selected = client.put(f"/api/sessions/{sid}/loan", json={"loan_id": "L16000"})
        assert selected.status_code == 200
        assert selected.json()["selected_loan_id"] == "L16000"
        assert selected.json()["selected_account_id"] == "A16000"
        scoped = selected.json()["preflight"]
        assert scoped["scoped_loan_id"] == "L16000"
        assert scoped["loan_count"] == scoped["account_count"] == 1
        assert "Unsupported_Source_Field__c" in scoped["unmapped_attributes"]
        assert "Credit report date" in scoped["unavailable_inputs"]
        assert "UNMAPPED_RELATION:Indicator:analysis_id" in scoped["mapping_warnings"]
        assert any(issue.startswith("CONFLICTING_ALIASES:CreditReport:fico") for issue in scoped["priority_issues"])
        assert graph_builder.call_count == 1
        assert all(len(doc.metadata["records"]) == 1 for doc in graph_builder.call_args.args[0])
        found = client.get(f"/api/sessions/{sid}/loans", params={"q": "L16000"}).json()
        assert found["results"] == [{"loan_id": "L16000", "account_id": "A16000",
                                    "account_name": "Synthetic client 16000"}]
        assert selected.json()["report"] is None
        assert len(selected.content) < 80_000
        timings["search_select_api"] = perf_counter() - started
        timings["source_reread"] = 0.0
        started = perf_counter()

        store = client.app.state.store
        saved = store.get(sid)
        assert saved is not None and saved.selected_loan_id == "L16000"
        assert saved.preflight is not None and saved.preflight.loan_count == 1
        for query, doc, payload in zip((1, 2, 4), original_documents, payloads):
            path = store.root / "uploads" / sid / doc["id"]
            assert path.read_bytes() == payload  # Every source byte, not just indexed columns.
            extracted = Document(name=doc["name"], extension=".xlsx", size=len(payload))
            extract(ExtractionAgent(), extracted, path, include_text=False)
            assert extracted.metadata["data_row_count"] == ROW_COUNT
            records = extracted.metadata["records"]
            assert len(records) == ROW_COUNT
            assert records[-1]["row"] == ROW_COUNT + 1
            assert records[-1]["values"]["Unsupported_Source_Field__c"] == f"raw-query-{query}-row-16000"
        timings["source_reread"] = perf_counter() - started
        provider.assert_not_called()
        submit.assert_not_called()
    print("\nSynthetic 3 x 16001 rows (seconds):", {key: round(value, 3) for key, value in timings.items()})