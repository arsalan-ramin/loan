"""V4 API integration: synthetic workbook evidence, mocked provider, no live calls.

New reviews intentionally do not run legacy evaluate_session: source rows are
not loan profiles, and v4 citations do not inherit the older PDF's p.4 rules.
"""

import csv
import hashlib
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.agents import DecisionAgent, RulesAgent
from backend.llm import AnthropicLLM, LLMError, report_context
from backend.main import create_app, process_session
from backend.models import CreateSession, Document, Report, Session
from backend.reports import render_report
from backend.tests.test_underwriting import document, rows
from backend.underwriting_engine import evaluate_underwriting
from backend.underwriting_models import IMPLEMENTATION_GUIDE_SHA256, SOURCE_DOCUMENT, SOURCE_SHA256


def commentary(context):
    trace = context["underwriting"]
    return {
        "recommended_route": "MANUAL_REVIEW", "decision_recommendation": None, "confidence": None,
        "positive_factors": [], "risk_factors": ["Unresolved policy conflicts"],
        "triggered_rules": [], "missing_information": [],
        "explanation": "Aggregate commentary only. Authorized human review required.",
        "requires_human_approval": True,
        "loan_ids": [loan["loan_id"] for loan in trace["loan_results"]],
        "evidence_refs": trace["evidence_refs"][:1],
    }


@pytest.fixture(autouse=True)
def provider(monkeypatch):
    calls = []

    def complete(_self, system, prompt):
        context = json.loads(prompt)
        calls.append(context)
        return json.dumps(commentary(context))

    monkeypatch.setattr(AnthropicLLM, "complete", complete)
    return calls


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / "data")) as active:
        yield active


def workbook(records, empty=False):
    book = Workbook()
    book.remove(book.active)
    for table, fields in records:
        sheet = book.create_sheet(table)
        sheet.append(list(fields))
        if not empty:
            sheet.append(list(fields.values()))
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def upload_three(client, mode=None, empty=False):
    response = client.post("/api/sessions", json={} if mode is None else {"mode": mode})
    assert response.status_code == 201
    sid = response.json()["id"]
    source = rows()
    source[0][1]["Name"] = "PRIVATE_RAW_WORKBOOK_SENTINEL"
    source[1][1]["requested_amount"] = "100000.01"
    groups = [(1, source[:2]), (2, source[3:]), (4, source[2:3])]
    response = client.post(f"/api/sessions/{sid}/documents", files=[
        ("files", (f"MTB_Query_{query}.xlsx", workbook(records, empty))) for query, records in groups
    ])
    assert response.status_code == 200
    validation = client.post(f"/api/sessions/{sid}/preflight")
    assert validation.status_code == 200
    accounts = client.get(f"/api/sessions/{sid}/accounts").json()
    if accounts["total"]:
        selected = client.put(f"/api/sessions/{sid}/account", json={"account_id": accounts["results"][0]["account_id"]})
        assert selected.status_code == 200
    return sid


def finish(client, sid):
    response = client.post(f"/api/sessions/{sid}/process")  # unchanged: no body
    assert response.status_code == 202
    client.app.state.executor.shutdown(wait=True)
    result = client.get(f"/api/sessions/{sid}").json()
    assert result["status"] == "completed", result["error"]
    return result


def test_mode_defaults_validation_and_not_ready(client):
    assert Session().mode == CreateSession().mode == "PRODUCTION"
    assert Report(decision="needs_review", summary="legacy", findings=[], document_count=0).underwriting is None
    sid = client.post("/api/sessions", json={}).json()["id"]
    assert client.get(f"/api/sessions/{sid}").json()["mode"] == "PRODUCTION"
    assert client.get(f"/api/sessions/{sid}/underwriting").status_code == 409
    assert client.get("/api/sessions/missing/underwriting").status_code == 404
    for bad in ("test", "APPROVE", "", None, 1):
        assert client.post("/api/sessions", json={"mode": bad}).status_code == 422
        assert client.put(f"/api/sessions/{sid}/mode", json={"mode": bad}).status_code == 422
    for extra in ("policy_config", "authority", "approved", "test_policy"):
        assert client.put(f"/api/sessions/{sid}/mode", json={"mode": "TEST", extra: {}}).status_code == 422
        assert client.post("/api/sessions", json={"mode": "TEST", extra: {}}).status_code == 422
    assert client.put(f"/api/sessions/{sid}/mode", json={}).status_code == 422


def test_production_real_upload_exact_json_csv_html_and_trace(client, provider, caplog):
    source = Path(__file__).parents[1] / "data" / SOURCE_DOCUMENT
    assert hashlib.sha256(source.read_bytes()).hexdigest() == SOURCE_SHA256
    sid = upload_three(client)
    with caplog.at_level("INFO", logger="backend.agents"):
        session = finish(client, sid)
    report = session["report"]
    batch = report["underwriting"]
    assert batch["source_available"] is True
    assert batch["implementation_guide_available"] is True
    assert batch["implementation_guide_source_sha256"] == IMPLEMENTATION_GUIDE_SHA256
    assert batch["expanded_spec_available"] is True
    assert batch["mode"] == "PRODUCTION"
    assert report["decision"] == "needs_review"
    loan, = batch["loan_results"]
    assert loan["loan_id"] == "L1"
    assert loan["final_action"] == "MANUAL_REVIEW"
    assert loan["recommendation_path"] == "POLICY_REVIEW_REQUIRED"
    assert loan["normalized_amounts_usd"]["loan_requested_amount_usd"] == "100000.01"
    assert loan["normalized_amounts_usd"]["total_group_credit_usd"] == "250000"
    assert {"REQUEST_THRESHOLD_100M_VS_100K", "INCREASE_WAIT_6_VS_12_MONTHS"} <= set(loan["conflicts"])
    assert report["policy_evaluation"]["profiles"][0]["profile_id"] == "L1"
    assert report["policy_evaluation"]["source_sha256"] == SOURCE_SHA256
    assert all(not rule["rule_id"].startswith("MTB-") for rule in report["policy_evaluation"]["rule_results"])
    assert all(rule["page"] in (None, 2) for rule in report["policy_evaluation"]["rule_results"])
    download = client.get(f"/api/sessions/{sid}/underwriting")
    assert download.json() == batch
    exported_loan = download.json()["loan_results"][0]
    guide_audit = exported_loan["normalized_data"]["implementation_guide_diagnostics"]["audit"]
    assert guide_audit["source_sha256"] == IMPLEMENTATION_GUIDE_SHA256
    assert len(guide_audit["input_fingerprint_sha256"]) == 64
    assert '"loan_requested_amount_usd":"100000.01"' in download.text
    assert download.headers["x-content-type-options"] == "nosniff"
    assert download.headers["content-disposition"].endswith('.json"')
    csv_response = client.get(f"/api/sessions/{sid}/underwriting?format=csv")
    csv_reader = csv.DictReader(io.StringIO(csv_response.text))
    required_csv_fields = {
        "loan_id", "borrower_account_id", "currency", "normalized_amounts_usd", "mode",
        "policy_version", "rule_results", "source_evidence", "missing_inputs", "conflicts",
        "recommendation_path", "eligibility_route", "final_action", "human_review_required", "explanation",
        "credit_assessment", "authorization", "strategy_branch", "implementation_guide_diagnostics",
        "implementation_guide_source_document", "implementation_guide_source_sha256", "implementation_guide_available",
    }
    assert required_csv_fields <= set(csv_reader.fieldnames or [])
    table = list(csv_reader)
    assert json.loads(table[0]["expanded_spec"]) == loan["expanded_spec"]
    assert len(loan["expanded_spec"]["decline_reasons"]) == 43
    assert json.loads(table[0]["normalized_amounts_usd"])["loan_requested_amount_usd"] == "100000.01"
    assert json.loads(table[0]["implementation_guide_diagnostics"]) == loan["normalized_data"]["implementation_guide_diagnostics"]
    for key in ("eligibility_route", "credit_assessment", "authorization", "strategy_branch"):
        assert table[0][key] == (loan[key] or "")
    assert client.get(f"/api/sessions/{sid}/underwriting?format=xml").status_code == 422
    html = client.get(f"/api/sessions/{sid}/report").text
    for value in ("V4 underwriting", SOURCE_DOCUMENT, SOURCE_SHA256, "USD 100000.01", "PDF", "p.2",
                  "POLICY_REVIEW_REQUIRED", "MANUAL_REVIEW", "Source evidence", "Conflicts", "Missing inputs",
                  "not LLM-generated", loan["source_evidence"][0]["id"]):
        assert value in html
    for value in (batch["implementation_guide_source_document"], IMPLEMENTATION_GUIDE_SHA256,
                  "Eligibility route", "Credit assessment candidate", "Authorization", "NOT_ISSUED",
                  "stage_0_schema_identity", "stage_6_explanation_and_audit", "input_fingerprint_sha256",
                  "NOT_EVALUATED"):
        assert value in html
    assert "<script" not in html
    assert "Complete decline-reason inventory" in html
    assert "Expanded rule registry and audit" in html
    assert len(provider) == 1
    assert "PRIVATE_RAW_WORKBOOK_SENTINEL" not in json.dumps(provider)
    assert "source_evidence" not in provider[0]["underwriting"]
    assert "underwriting_evaluation mode=PRODUCTION" in caplog.text
    assert "PRIVATE_RAW_WORKBOOK_SENTINEL" not in caplog.text
    assert not any("L1" in record.message for record in caplog.records)
    assert "v4-USD-verified-draft-2" in session["steps"][3]["detail"]


def test_mode_change_invalidates_and_processing_locks(client):
    sid = upload_three(client)
    finish(client, sid)
    changed = client.put(f"/api/sessions/{sid}/mode", json={"mode": "TEST"}).json()
    assert changed["mode"] == "TEST"
    assert changed["report"] is None
    assert changed["status"] == "ready"
    assert all(step["status"] == "pending" for step in changed["steps"])
    assert client.get(f"/api/sessions/{sid}/underwriting").status_code == 409
    store = client.app.state.store
    session = store.get(sid)
    session.status = "processing"
    store.save(session)
    assert client.put(f"/api/sessions/{sid}/mode", json={"mode": "PRODUCTION"}).status_code == 409
    assert store.get(sid).mode == "TEST"
    process_session(store, sid)
    loan = client.get(f"/api/sessions/{sid}/underwriting").json()["loan_results"][0]
    assert loan["recommendation_path"] == "SIMULATED_REVIEW"
    assert loan["final_action"] == "NO_OPERATIONAL_ACTION"
    assert {"REQUEST_THRESHOLD_100M_VS_100K", "INCREASE_WAIT_6_VS_12_MONTHS"} <= set(loan["conflicts"])


def test_test_mode_never_silently_enables_synthetic_approval(client):
    sid = upload_three(client, mode="TEST")
    loan = finish(client, sid)["report"]["underwriting"]["loan_results"][0]
    assert loan["mode"] == "TEST"
    assert loan["recommendation_path"] == "SIMULATED_REVIEW"
    assert loan["final_action"] == "NO_OPERATIONAL_ACTION"


def test_empty_templates_and_unidentified_rows_never_fabricate_loans(client):
    sid = upload_three(client, empty=True)
    report = finish(client, sid)["report"]
    assert report["underwriting"]["loan_results"] == []
    assert "NO_EXPLICIT_LOAN_RECORDS" in report["underwriting"]["data_quality_issues"]
    assert report["policy_evaluation"]["profiles"] == []
    assert report["policy_evaluation"]["rule_results"] == []
    assert "No explicit loan records" in client.get(f"/api/sessions/{sid}/report").text
    unidentified = Session(documents=[Document(name="MTB_Query_4.xlsx", extension=".xlsx", size=1,
        metadata={"records": [{"sheet": "CreditReport", "row": 2, "values": {"credit_freeze": True}}]})])
    result = DecisionAgent().decide(unidentified, RulesAgent().evaluate(unidentified))
    assert result.underwriting.loan_results == result.policy_evaluation.profiles == []
    assert not any(rule.rule_id == "MTB-CREDIT-FREEZE" for rule in result.policy_evaluation.rule_results)


@pytest.mark.parametrize("change", [
    {"recommended_route": "APPROVE"}, {"decision_recommendation": "DENY"},
    {"loan_ids": ["UNKNOWN"]}, {"triggered_rules": ["INVENTED"]}, {"evidence_refs": ["ev-unknown"]},
    {"confidence": 0.9}, {"final_action": "APPROVE"}, {"recommendation_path": "SIMULATED_APPROVE_PATH"},
])
def test_invalid_commentary_never_loses_or_mutates_results(monkeypatch, change):
    session = Session(documents=[Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=1,
                                         metadata=document(rows()).metadata)])
    expected = evaluate_underwriting(session.documents).model_dump_json()

    def invalid(_self, _system, prompt):
        return json.dumps({**commentary(json.loads(prompt)), **change})

    monkeypatch.setattr(AnthropicLLM, "complete", invalid)
    report = DecisionAgent().decide(session, [])
    assert report.underwriting.model_dump_json() == expected
    assert report.decision == "needs_review"
    assert report.llm_error
    assert report.llm_analysis is report.llm_policy_output is None


def test_provider_failure_preserves_all_results(client, monkeypatch):
    def failure(*_args):
        raise LLMError("Mock provider unavailable", 503)

    monkeypatch.setattr(AnthropicLLM, "complete", failure)
    sid = upload_three(client)
    report = finish(client, sid)["report"]
    assert report["llm_error"] == "Mock provider unavailable"
    assert report["underwriting"]["loan_results"][0]["final_action"] == "MANUAL_REVIEW"
    assert client.get(f"/api/sessions/{sid}/underwriting").status_code == 200
    assert "Mock provider unavailable" in client.get(f"/api/sessions/{sid}/report").text


def test_bounded_one_call_for_many_loans_and_deterministic_explanations(provider):
    source = rows()
    source.extend(("Loan", {**source[1][1], "Id": f"L{index:03}"}) for index in range(2, 32))
    session = Session(documents=[Document(name="MTB_Query_1.xlsx", extension=".xlsx", size=1,
                                         metadata=document(source).metadata)])
    report = DecisionAgent().decide(session, [])
    assert len(report.underwriting.loan_results) == 31
    assert len(provider) == 1
    trace = report_context(report)["underwriting"]
    assert len(trace["loan_results"]) == 20
    assert sum(len(loan["rule_results"]) for loan in trace["loan_results"]) == 40
    assert len(trace["evidence_refs"]) <= 40
    assert sum(len(loan["expanded_spec"]["reason_diagnostics"]) for loan in trace["loan_results"]) <= 40
    assert trace["loan_count"] == 31
    assert all("Authorized human review" in loan.explanation for loan in report.underwriting.loan_results)
    assert all("normalized_data" not in loan for loan in trace["loan_results"])
    assert report.llm_analysis == "Aggregate commentary only. Authorized human review required."


def test_csv_formula_sanitization_and_html_escaping(client):
    sid = upload_three(client)
    finish(client, sid)
    store = client.app.state.store
    session = store.get(sid)
    loan = session.report.underwriting.loan_results[0]
    # Exercise export hardening independently of the graph's strict ID grammar.
    loan.loan_id = " \t=HYPERLINK(1)"
    loan.explanation = "@SUM(1,2) <script>alert(1)</script>"
    loan.missing_inputs.append("<img src=x onerror=alert(1)>")
    loan.normalized_data["implementation_guide_diagnostics"]["stages"]["stage_6_explanation_and_audit"] = "<script>untrusted</script>"
    store.save(session)
    response = client.get(f"/api/sessions/{sid}/underwriting?format=csv")
    record, = csv.DictReader(io.StringIO(response.text))
    assert record["loan_id"].startswith("'")
    assert record["explanation"].startswith("'@")
    html = render_report(store.get(sid))
    assert "<script" not in html and "<img" not in html
    assert "&lt;script&gt;" in html and "&lt;img" in html


def test_mode_and_exact_batch_survive_restart(tmp_path):
    root = tmp_path / "persistent"
    with TestClient(create_app(root)) as client:
        sid = upload_three(client, mode="TEST")
        expected = finish(client, sid)["report"]["underwriting"]
    with TestClient(create_app(root)) as client:
        assert client.get(f"/api/sessions/{sid}").json()["mode"] == "TEST"
        assert client.get(f"/api/sessions/{sid}/underwriting").json() == expected