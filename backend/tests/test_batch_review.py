"""Whole stored-dataset batch review: deterministic status, Claude rationale, latest run only."""

import json
import threading
import time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.batch_review import BatchReviewer, readable_rationale
from backend.llm import AnthropicLLM, LLMError
from backend.loan_dataset import LoanDataset
from backend.main import create_app
from backend.tests.test_loan_dataset import upload_set, wait_ready, write_set


def claude_output(context, explanation=None):
    loan = context["underwriting"]["loan_results"][0]
    label = loan["demo_decision"]["label"]
    return json.dumps({
        "recommended_route": "MANUAL_REVIEW", "decision_recommendation": None, "confidence": None,
        "requires_human_approval": True, "loan_ids": [loan["loan_id"]],
        "positive_factors": [], "risk_factors": [], "missing_information": [],
        "explanation": explanation or f"Result: {loan['loan_id']} is {label}. Why: decision-flow criteria.",
        "triggered_rules": [], "evidence_refs": [], "check_reconciliation": [],
    })


@pytest.fixture
def claude(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    calls = []

    def complete(self, system, prompt):
        context = json.loads(prompt)
        calls.append(context["underwriting"]["loan_results"][0]["loan_id"])
        return claude_output(context)

    monkeypatch.setattr(AnthropicLLM, "complete", complete)
    return calls


@pytest.fixture
def reviewer(tmp_path):
    dataset = LoanDataset(tmp_path / "data")
    _, files = write_set(tmp_path / "files", count=8)
    assert dataset.replace(files)
    return BatchReviewer(dataset, concurrency=3)


def wait_done(reviewer):
    for _ in range(400):
        if not reviewer.running():
            return reviewer.status()
        time.sleep(0.02)
    pytest.fail("Batch review did not finish")


def test_batch_reviews_every_loan_with_deterministic_status_and_claude_rationale(reviewer, claude):
    state = reviewer.start()
    assert state["status"] == "running" and state["total"] == 8
    state = wait_done(reviewer)
    assert state["status"] == "completed" and state["processed"] == 8 and state["errors"] == 0
    assert state["approved"] + state["declined"] + state["referred"] == 8 and state["dataset_current"]
    assert sorted(claude) == [f"L{i:05d}" for i in range(8)]
    page = reviewer.results(limit=5)
    assert page["total"] == 8 and [row["loan_id"] for row in page["results"]] == [f"L{i:05d}" for i in range(5)]
    row = page["results"][0]
    assert row["account_id"] == "A00000" and row["status"] in {"Approved", "Decline", "Refer"}
    assert row["rationale"].startswith("Result: L00000") and "\nWhy:" in row["rationale"]
    assert row["engine_rationale"] and row["error"] is None
    assert reviewer.results(status=row["status"])["total"] >= 1


def test_export_has_requested_columns_and_sanitizes_formulas(reviewer, claude, monkeypatch):
    monkeypatch.setattr(AnthropicLLM, "complete",
                        lambda self, system, prompt: claude_output(json.loads(prompt), "=HYPERLINK(\"x\") Result: x"))
    with pytest.raises(Exception, match="no batch results"):
        reviewer.export_xlsx()
    reviewer.start()
    wait_done(reviewer)
    sheet = load_workbook(BytesIO(reviewer.export_xlsx())).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Loan ID", "Account ID", "Status", "Rationale / result")
    assert len(rows) == 9 and rows[1][:2] == ("L00000", "A00000")
    assert rows[1][2] in {"Approved", "Decline", "Refer"}
    assert rows[1][3].startswith("'=HYPERLINK")


def test_per_loan_claude_failure_is_recorded_and_resume_retries_only_failures(reviewer, claude, monkeypatch):
    original = AnthropicLLM.complete

    def flaky(self, system, prompt):
        if json.loads(prompt)["underwriting"]["loan_results"][0]["loan_id"] == "L00003":
            raise LLMError("Anthropic request failed.", 502)
        return original(self, system, prompt)

    monkeypatch.setattr(AnthropicLLM, "complete", flaky)
    reviewer.start()
    state = wait_done(reviewer)
    assert state["status"] == "completed" and state["processed"] == 8 and state["errors"] == 1
    failed = reviewer.results(status="Error")["results"][0]
    assert failed["loan_id"] == "L00003" and failed["status"] in {"Approved", "Decline", "Refer"}
    assert failed["rationale"] is None and "Claude rationale unavailable" in failed["error"]

    monkeypatch.setattr(AnthropicLLM, "complete", original)
    claude.clear()
    reviewer.start(resume=True)
    state = wait_done(reviewer)
    assert claude == ["L00003"]
    assert state["status"] == "completed" and state["processed"] == 8 and state["errors"] == 0


def test_retry_selected_failed_loans_reprocesses_only_those(reviewer, claude, monkeypatch):
    original = AnthropicLLM.complete
    failing = {"L00001", "L00004", "L00006"}

    def flaky(self, system, prompt):
        if json.loads(prompt)["underwriting"]["loan_results"][0]["loan_id"] in failing:
            raise LLMError("Anthropic request failed.", 502)
        return original(self, system, prompt)

    monkeypatch.setattr(AnthropicLLM, "complete", flaky)
    reviewer.start()
    assert wait_done(reviewer)["errors"] == 3
    rows = {row["loan_id"]: row["processing"] for row in reviewer.results()["results"]}
    assert rows["L00001"] == "failed" and rows["L00000"] == "succeeded"
    assert reviewer.results(status="Succeeded")["total"] == 5

    monkeypatch.setattr(AnthropicLLM, "complete", original)
    claude.clear()
    with pytest.raises(Exception, match="None of the selected loans"):
        reviewer.start(loan_ids=["L00000", "NOPE"])
    state = reviewer.start(loan_ids=["L00004", "L00006", "L00000"])
    assert state["job_mode"] == "retry" and state["job_total"] == 2
    assert state["retrying_loan_ids"] == ["L00004", "L00006"]
    state = wait_done(reviewer)
    assert sorted(claude) == ["L00004", "L00006"]
    assert state["status"] == "completed" and state["processed"] == 8 and state["errors"] == 1
    assert [row["loan_id"] for row in reviewer.results(status="Error")["results"]] == ["L00001"]
    # Retried rows keep their original ordering.
    assert [row["loan_id"] for row in reviewer.results()["results"]] == [f"L{i:05d}" for i in range(8)]

    claude.clear()
    reviewer.start(failed_only=True)
    assert wait_done(reviewer)["errors"] == 0 and claude == ["L00001"]
    with pytest.raises(Exception, match="no failed loans"):
        reviewer.start(failed_only=True)


def test_provider_configuration_failure_stops_the_run(reviewer, claude, monkeypatch):
    def unavailable(self, system, prompt):
        raise LLMError("Anthropic authentication failed. Check the backend ANTHROPIC_API_KEY.", 503)

    monkeypatch.setattr(AnthropicLLM, "complete", unavailable)
    reviewer.start()
    state = wait_done(reviewer)
    assert state["status"] == "failed" and "authentication failed" in state["error"]
    assert state["processed"] == 0


def test_missing_api_key_is_rejected_before_any_work(reviewer, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    with pytest.raises(Exception, match="not configured") as error:
        reviewer.start()
    assert error.value.status_code == 503 and reviewer.status()["status"] == "none"


def test_cancel_keeps_finished_results_and_new_run_replaces_latest(reviewer, claude, monkeypatch):
    gate, started, original = threading.Event(), threading.Event(), AnthropicLLM.complete

    def slow(self, system, prompt):
        started.set()
        gate.wait(5)
        return original(self, system, prompt)

    monkeypatch.setattr(AnthropicLLM, "complete", slow)
    reviewer.start()
    with pytest.raises(Exception, match="already running"):
        reviewer.start()
    assert started.wait(10)
    reviewer.cancel()
    gate.set()
    state = wait_done(reviewer)
    assert state["status"] == "cancelled" and 0 < state["processed"] < 8

    monkeypatch.setattr(AnthropicLLM, "complete", original)
    reviewer.start()
    assert wait_done(reviewer)["processed"] == 8


def test_replacing_stored_data_clears_batch_results(reviewer, claude, tmp_path):
    reviewer.start()
    wait_done(reviewer)
    _, files = write_set(tmp_path / "other", count=3)
    assert reviewer.dataset.replace(files)
    state = reviewer.status()
    assert state["status"] == "none" and state["processed"] == 0 and state["dataset_loan_count"] == 3
    with pytest.raises(Exception, match="nothing|no batch review"):
        reviewer.start(resume=True)


def test_interrupted_run_is_recoverable(reviewer, claude):
    with reviewer.dataset.connect() as conn:
        conn.execute("INSERT INTO batch_run VALUES (1, 'running', 'x', NULL, 8, "
                     "(SELECT value FROM dataset_meta WHERE key='updated_at'), 'm', NULL)")
    reviewer.recover_interrupted()
    assert reviewer.status()["status"] == "interrupted"
    reviewer.start(resume=True)
    assert wait_done(reviewer)["processed"] == 8


def test_readable_rationale_puts_sections_on_lines():
    assert readable_rationale("Result: L1 is Refer. Why: x. What is unresolved: y.") == \
        "Result: L1 is Refer.\nWhy: x.\nWhat is unresolved: y."


def test_batch_api_flow(tmp_path, claude):
    with TestClient(create_app(tmp_path / "data")) as client:
        assert client.get("/api/batch").json()["status"] == "none"
        assert client.post("/api/batch", json={}).status_code == 409
        upload_set(client, 5)
        wait_ready(client, 5)
        started = client.post("/api/batch", json={})
        assert started.status_code == 202, started.text
        for _ in range(400):
            state = client.get("/api/batch").json()
            if state["status"] != "running":
                break
            time.sleep(0.02)
        assert state["status"] == "completed" and state["processed"] == 5
        page = client.get("/api/batch/results", params={"limit": 2}).json()
        assert page["total"] == 5 and len(page["results"]) == 2
        assert client.get("/api/batch/results", params={"status": "Bogus"}).status_code == 422
        exported = client.get("/api/batch/export")
        assert exported.status_code == 200
        assert "attachment" in exported.headers["content-disposition"]
        assert len(list(load_workbook(BytesIO(exported.content)).active.iter_rows())) == 6
        assert client.post("/api/batch/cancel").status_code == 409
        assert client.post("/api/batch", json={"extra": 1}).status_code == 422
        assert client.post("/api/batch", json={"loan_ids": []}).status_code == 422
        assert client.post("/api/batch", json={"loan_ids": ["L00001"]}).status_code == 409
        assert client.post("/api/batch", json={"failed_only": True}).status_code == 409
